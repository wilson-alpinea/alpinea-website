import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from PIL import Image

class MarcadorXlsx:
    def __init__(self, caminho, aba):
        self.wb = openpyxl.load_workbook(caminho)
        self.ws = self.wb[aba]
        self.usados = set()
        self.tops = {(m.min_row, m.min_col): m for m in self.ws.merged_cells.ranges}
        self.cobertas = {(r, c) for m in self.ws.merged_cells.ranges for r in range(m.min_row, m.max_row + 1)
                         for c in range(m.min_col, m.max_col + 1) if (r, c) != (m.min_row, m.min_col)}

    def t(self, coord, chave, permitir_texto=False):
        assert chave not in self.usados, chave
        col, row = coordinate_from_string(coord)
        c = column_index_from_string(col)
        assert (row, c) not in self.cobertas, f'{coord} é célula coberta por mesclagem'
        atual = self.ws[coord].value
        assert permitir_texto or atual in (None, ''), f'{coord} já tem {atual!r}'
        self.ws[coord].value = '{{%s}}' % chave
        self.usados.add(chave)

    def foto(self, faixa, arquivo_jpg):
        """Imagem-placeholder (JPEG branco) ancorada na faixa; em produção o
        arquivo xl/media/image1.jpeg é trocado pela foto do candidato."""
        ini, fim = faixa.split(':')
        ws = self.ws
        c0, r0 = coordinate_from_string(ini); c1, r1 = coordinate_from_string(fim)
        c0, c1 = column_index_from_string(c0), column_index_from_string(c1)
        def largura_col(c):
            for d in ws.column_dimensions.values():
                if d.min and d.max and d.min <= c <= d.max and d.width:
                    return d.width
            return ws.sheet_format.defaultColWidth or 8.43
        larg = sum(largura_col(c) * 7 + 5 for c in range(c0, c1 + 1))
        alt = sum((ws.row_dimensions[r].height or ws.sheet_format.defaultRowHeight or 15) for r in range(r0, r1 + 1)) * 96 / 72
        w, h = int(larg * 0.86), int(alt * 0.9)
        Image.new('RGB', (w, h), 'white').save(arquivo_jpg, quality=60)
        img = XLImage(arquivo_jpg); img.width, img.height = w, h
        # centraliza com um pequeno recuo
        from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        dx, dy = int((larg - w) / 2), int((alt - h) / 2)
        img.anchor = OneCellAnchor(_from=AnchorMarker(col=c0 - 1, colOff=pixels_to_EMU(dx), row=r0 - 1, rowOff=pixels_to_EMU(dy)),
                                   ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
        ws.add_image(img)
        return w, h

    def salvar(self, caminho):
        self.wb.save(caminho)
        return sorted(self.usados)
