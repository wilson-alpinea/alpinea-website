"""Marca um .docx original com {{tokens}} (texto) e tags cb:KEY (caixas de
seleção). O preenchimento em produção (lib/empregos/preencherModelo.ts) só
troca os tokens — o resto do arquivo fica idêntico ao original."""
import copy, docx
from docx.oxml.ns import qn
from lxml import etree
W14 = 'http://schemas.microsoft.com/office/word/2010/wordml'

class Marcador:
    def __init__(self, caminho):
        self.d = docx.Document(caminho)
        self.body = self.d.element.body
        self.cbs = [s for s in self.body.iter(qn('w:sdt')) if s.find('.//{%s}checkbox' % W14) is not None]
        self.usados = set()

    def cell(self, t, r, c):
        return self.d.tables[t].rows[r].cells[c]

    def _run(self, p, texto, modelo=None):
        novo = etree.Element(qn('w:r'))
        rpr = None
        if modelo is not None and modelo.find(qn('w:rPr')) is not None:
            rpr = copy.deepcopy(modelo.find(qn('w:rPr')))
        else:
            ppr = p.find(qn('w:pPr'))
            if ppr is not None and ppr.find(qn('w:rPr')) is not None:
                rpr = copy.deepcopy(ppr.find(qn('w:rPr')))
                rpr.tag = qn('w:rPr')
        if rpr is not None:
            for x in rpr.findall(qn('w:b')) + rpr.findall(qn('w:bCs')) + rpr.findall(qn('w:color')):
                rpr.remove(x)
            # valores em azul-escuro, para destacar do impresso
            cor = etree.SubElement(rpr, qn('w:color')); cor.set(qn('w:val'), '1F3864')
            novo.append(rpr)
        t = etree.SubElement(novo, qn('w:t')); t.text = texto
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        return novo

    def texto(self, t, r, c, chave, par=0):
        """Token num campo vazio (ou no fim do parágrafo `par` da célula)."""
        self._uso(chave)
        p = self.cell(t, r, c).paragraphs[par]._p
        runs = p.findall(qn('w:r'))
        novo = self._run(p, ('' if not runs else ' ') + '{{%s}}' % chave, runs[-1] if runs else None)
        p.append(novo)

    def depois(self, t, r, c, rotulo, chave):
        """Token logo depois do texto `rotulo` dentro da célula."""
        self._uso(chave)
        for par in self.cell(t, r, c).paragraphs:
            p = par._p
            acum = ''
            for el in p.iter(qn('w:t')):
                acum += el.text or ''
                if rotulo in acum:
                    run = el.getparent()
                    alvo = run
                    # se o run estiver dentro de sdt/hyperlink, sobe até filho do parágrafo
                    while alvo.getparent() is not p: alvo = alvo.getparent()
                    alvo.addnext(self._run(p, ' {{%s}}' % chave, run))
                    return
        raise ValueError(f'rótulo não achado: {rotulo!r} em {(t, r, c)}')

    def caixa(self, i, chave):
        self._uso('cb:' + chave)
        s = self.cbs[i]
        pr = s.find(qn('w:sdtPr'))
        tag = pr.find(qn('w:tag'))
        if tag is None:
            tag = etree.SubElement(pr, qn('w:tag'))
        tag.set(qn('w:val'), 'cb:' + chave)

    def _uso(self, k):
        assert k not in self.usados, k
        self.usados.add(k)

    def salvar(self, caminho):
        self.d.save(caminho)
        return sorted(self.usados)
